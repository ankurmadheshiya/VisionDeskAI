import os
import csv
from datetime import datetime

class SafetyReportGenerator:
    """
    Generates downloadable reports in CSV, Excel, and PDF formats for workplace safety data.
    """

    def __init__(self, output_dir: str = None):
        if output_dir is None:
            output_dir = os.path.join(os.path.dirname(__file__), "output")
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def generate_csv_report(self, violations_data: list, report_title: str = "Workplace_Safety_Report") -> str:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{report_title}_{timestamp}.csv"
        filepath = os.path.join(self.output_dir, filename)

        headers = ["ID", "Violation Type", "Severity", "Status", "Location", "Department", "Confidence", "Timestamp", "Details"]

        with open(filepath, mode="w", newline="", encoding="utf-8") as file:
            writer = csv.writer(file)
            writer.writerow(headers)
            for item in violations_data:
                writer.writerow([
                    item.get("id", ""),
                    item.get("violation_type", ""),
                    item.get("severity", ""),
                    item.get("status", ""),
                    item.get("location", ""),
                    item.get("department", ""),
                    item.get("confidence", ""),
                    item.get("timestamp", ""),
                    item.get("details", "")
                ])

        return filepath

    def generate_excel_report(self, violations_data: list, report_title: str = "Workplace_Safety_Report") -> str:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            # Fallback to CSV if openpyxl is unavailable
            return self.generate_csv_report(violations_data, report_title)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{report_title}_{timestamp}.xlsx"
        filepath = os.path.join(self.output_dir, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Safety Violations"

        # Title block
        ws.merge_cells("A1:I1")
        ws["A1"] = f"VisionDesk AI - {report_title.replace('_', ' ')}"
        ws["A1"].font = Font(name="Arial", size=16, bold=True, color="1F2937")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:I2")
        ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Confidential Safety Audit"
        ws["A2"].font = Font(name="Arial", size=10, italic=True, color="6B7280")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        # Headers
        headers = ["ID", "Violation Type", "Severity", "Status", "Location", "Department", "Confidence", "Timestamp", "Details"]
        ws.append([]) # Empty line
        ws.append(headers)

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data rows
        for row_idx, item in enumerate(violations_data, start=5):
            ws.append([
                item.get("id", ""),
                item.get("violation_type", ""),
                item.get("severity", ""),
                item.get("status", ""),
                item.get("location", ""),
                item.get("department", ""),
                f"{float(item.get('confidence', 0.9)) * 100:.1f}%",
                str(item.get("timestamp", "")),
                item.get("details", "")
            ])

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(filepath)
        return filepath

    def generate_pdf_report(self, violations_data: list, report_title: str = "Workplace_Safety_Report") -> str:
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib import colors
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        except ImportError:
            # Fallback to CSV if reportlab is unavailable
            return self.generate_csv_report(violations_data, report_title)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{report_title}_{timestamp}.pdf"
        filepath = os.path.join(self.output_dir, filename)

        doc = SimpleDocTemplate(filepath, pagesize=letter)
        story = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=6
        )
        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            parent=styles["Normal"],
            fontSize=10,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=15
        )

        story.append(Paragraph("VisionDesk AI - Safety & Compliance Report", title_style))
        story.append(Paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Violations Logged: {len(violations_data)}", subtitle_style))
        story.append(Spacer(1, 10))

        # Build table
        table_data = [["ID", "Violation Type", "Severity", "Status", "Location", "Department"]]
        for item in violations_data[:20]: # Limit for PDF clean page fit
            table_data.append([
                str(item.get("id", "")),
                str(item.get("violation_type", "")),
                str(item.get("severity", "")),
                str(item.get("status", "")),
                str(item.get("location", "")),
                str(item.get("department", ""))
            ])

        t = Table(table_data, colWidths=[30, 110, 70, 70, 110, 110])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0F172A")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor("#F8FAFC")),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
        ]))
        story.append(t)

        doc.build(story)
        return filepath
